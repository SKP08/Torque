"""Turn a file into plain text for searching. Local only; unreadable files are skipped, never fatal."""

import re
from pathlib import Path

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".rst", ".log", ".csv", ".tsv", ".json", ".yaml", ".yml", ".toml",
                 ".ini", ".cfg", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".cs", ".go", ".rs", ".sql",
                 ".sh", ".ps1", ".bat", ".html", ".htm", ".xml", ".tex"}
DOC_SUFFIXES = {".pdf", ".docx"}
SUPPORTED = TEXT_SUFFIXES | DOC_SUFFIXES

MAX_BYTES = 25 * 1024 * 1024
MAX_PDF_PAGES = 400

# Never index files that look like secrets, even inside a folder the user chose.
SECRET_NAME = re.compile(
    r"(?:^|[\\/])(?:\.env[^\\/]*|id_rsa[^\\/]*|id_ed25519[^\\/]*|[^\\/]*\.(?:pem|key|pfx|p12|kdbx|ppk|gpg|asc|keystore)|"
    r"[^\\/]*(?:password|passwd|secret|credential|token|private[_-]?key|api[_-]?key|wallet|seed[_-]?phrase)[^\\/]*|"
    r"known_hosts|authorized_keys|shadow|\.netrc|\.pgpass)$",
    re.IGNORECASE,
)
SKIP_DIRS = {".git", ".svn", "node_modules", "__pycache__", ".venv", "venv", "env", ".idea", ".vscode", "dist",
             "build", "site-packages", ".cache", "$recycle.bin", "appdata", ".ssh", ".aws", ".gnupg"}


def is_secret(path) -> bool:
    return bool(SECRET_NAME.search(str(path)))


def read_text(path) -> list:
    """[(location, text)] - location is 'page 3' for PDFs, '' otherwise. Empty when unreadable."""

    path = Path(path)
    suffix = path.suffix.lower()

    try:
        if suffix not in SUPPORTED or path.stat().st_size > MAX_BYTES or is_secret(path):
            return []

        if suffix == ".pdf":
            return _pdf(path)

        if suffix == ".docx":
            return _docx(path)

        text = path.read_text(encoding="utf-8", errors="replace")

        if suffix in (".html", ".htm", ".xml"):
            text = re.sub(r"<(script|style)\b.*?</\1>", " ", text, flags=re.S | re.I)
            text = re.sub(r"<[^>]+>", " ", text)

        return [("", text)]

    except Exception:  # noqa: BLE001 - a bad file must never stop indexing
        return []


def _pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(str(path))

    if reader.is_encrypted:
        return []

    pages = []

    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
        text = page.extract_text() or ""

        if text.strip():
            pages.append((f"page {number}", text))

    return pages


def _docx(path):
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))

    return [("", "\n".join(parts))]
